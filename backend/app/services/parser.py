import pdfplumber
import docx
import os
import base64
import io
from typing import Dict, Any

def parse_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    
    # Also grab tables
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text)
            full_text.append(" | ".join(row_data))
            
    return "\n".join(full_text)

def parse_pdf(file_path: str) -> Dict[str, Any]:
    text = ""
    images_base64 = []
    
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
            
        # If text is too short, we assume it's a scanned PDF and extract page images for OpenAI Vision
        if len(text.strip()) < 50:
            for page in pdf.pages:
                im = page.to_image(resolution=200)
                buffered = io.BytesIO()
                im.original.save(buffered, format="JPEG")
                img_str = base64.b64encode(buffered.getvalue()).decode("utf-8")
                images_base64.append(f"data:image/jpeg;base64,{img_str}")
                
    return {
        "text": text.strip(),
        "images": images_base64,
        "is_scanned": len(text.strip()) < 50
    }

def extract_cv_content(file_path: str, filename: str) -> Dict[str, Any]:
    """
    Extracts text from PDF or DOCX. 
    If PDF text extraction yields too little text (scanned PDF),
    it extracts base64 images to be sent to OpenAI Vision API.
    """
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".docx":
        return {
            "text": parse_docx(file_path),
            "images": [],
            "is_scanned": False
        }
    elif ext == ".pdf":
        return parse_pdf(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")
